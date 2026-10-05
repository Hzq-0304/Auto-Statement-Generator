"""仅在本机原始样例存在时执行；CI 不需要上传用户数据。"""
import hashlib
from pathlib import Path
import tempfile
import unittest
import openpyxl
from app.parser import import_excel
from app.matching import reconcile
from app.exporter import export_statement

ROOT=Path(__file__).resolve().parents[1]
QUOTES=list(ROOT.glob('*报价*.xlsx'));DELIVERIES=list(ROOT.glob('送货*.xlsx'))

@unittest.skipUnless(QUOTES and DELIVERIES,'本地业务样例未提交到仓库')
class SampleTests(unittest.TestCase):
    def test_full_roundtrip(self):
        paths=[QUOTES[0],DELIVERIES[0]]
        before=[hashlib.sha256(p.read_bytes()).hexdigest() for p in paths]
        q=import_excel(paths[0],'quote');d=import_excel(paths[1],'delivery');m=reconcile(q.items,d.items)
        self.assertEqual(len(q.items),327);self.assertEqual(len(d.items),99)
        self.assertEqual(sum(i.price is None for i in q.items),7)
        self.assertTrue(any('正文日期' in i.message for i in d.issues))
        self.assertTrue(any('单号出现在多张表' in i.message for i in d.issues))
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'样例结果.xlsx';total=export_statement(path,q,d,m)
            ws=openpyxl.load_workbook(path,data_only=True).active
            self.assertEqual(ws['K7'].value,'1');self.assertEqual(ws['C7'].value,'2026080301')
            self.assertAlmostEqual(ws['H107'].value,float(total))
            self.assertEqual(ws['D105'].value,d.items[-1].name)
        self.assertEqual(before,[hashlib.sha256(p.read_bytes()).hexdigest() for p in paths])

if __name__=='__main__':unittest.main()
