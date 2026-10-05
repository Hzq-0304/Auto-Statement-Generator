"""导入草稿的真实Excel回归：修正、强制提交、金额保护与进度恢复。"""
import hashlib
import json
from datetime import date
from decimal import Decimal as D
from pathlib import Path
import tempfile
import unittest
import openpyxl
from app.parser import import_excel
from app.models import ImportFailure
from app.review import edit_item, commit_review, force_warning
from app.matching import reconcile, confirm
from app.exporter import export_statement
from app.session import save_session, load_session


class ReviewTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def book(self, name, headers, rows, meta=''):
        wb = openpyxl.Workbook(); ws = wb.active
        ws['A1'] = meta
        for c, h in enumerate(headers, 3): ws.cell(4, c, h)
        for r, values in enumerate(rows, 5):
            for c, v in enumerate(values, 3): ws.cell(r, c, v)
        p = self.root / name; wb.save(p); wb.close()
        return p

    def test_bad_rows_preserved_then_corrected_without_changing_source(self):
        p = self.book('d.xlsx', ['序号','品名','数量','单价','金额'],
                      [['D1',None,'2箱',3,6], ['D2','螺母',None,2,8]])
        original = hashlib.sha256(p.read_bytes()).digest()
        with self.assertRaises(ImportFailure): import_excel(p, 'delivery')
        draft = import_excel(p, 'delivery', preview=True)
        self.assertEqual(len(draft.items), 2)
        self.assertEqual(draft.unresolved, 2)
        with self.assertRaisesRegex(ValueError,'未修正'): commit_review(draft)
        edit_item(draft,0,{'name':'螺丝','quantity':'2'})
        edit_item(draft,1,{'quantity':'4'})
        result = commit_review(draft)
        self.assertEqual(result.unresolved, 0)
        self.assertTrue(result.reviewed); self.assertFalse(result.forced)
        self.assertEqual(result.items[0].original_values['quantity'],'2箱')
        self.assertEqual(result.items[0].quantity,D(2))
        self.assertEqual(len(result.items[0].edit_history),2)
        edit_item(draft,0,{'quantity':'5'})
        self.assertEqual(result.items[0].quantity,D(2))  # 提交是独立快照
        self.assertEqual(original,hashlib.sha256(p.read_bytes()).digest())

    def test_nameless_amount_only_and_uncached_quantity_are_not_empty_templates(self):
        p=self.book('d.xlsx',['品名','数量','单价','金额'],[[None,None,3,6],[None,'=2+3',None,None]])
        d=import_excel(p,'delivery',preview=True)
        self.assertEqual(len(d.items),2)
        for item in d.items:
            self.assertIn('name',item.validation_errors)
            self.assertIn('quantity',item.validation_errors)

    def test_forced_import_keeps_bad_quantity_and_blocks_formal(self):
        qp=self.book('q.xlsx',['品名','单价'],[['螺丝',3]])
        dp=self.book('d.xlsx',['品名','数量'],[['螺丝','两箱'],['螺丝',2]])
        q=commit_review(import_excel(qp,'quote',preview=True))
        d=commit_review(import_excel(dp,'delivery',preview=True),force=True)
        self.assertTrue(d.forced);self.assertIn('1 条尚有错误',force_warning(d))
        matches=reconcile(q.items,d.items)
        self.assertFalse(matches[0].ready);self.assertTrue(matches[1].ready)
        with self.assertRaises(ValueError): confirm(matches[0],q.items[0],D(1),D(3),'测试')
        out=self.root/'draft.xlsx'
        self.assertEqual(export_statement(out,q,d,matches),D(6))
        wb=openpyxl.load_workbook(out,data_only=True)
        self.assertEqual(wb.active['F7'].value,'两箱');self.assertIsNone(wb.active['H7'].value)
        self.assertEqual(wb['来源明细']['V2'].value,'强制导入');wb.close()
        with self.assertRaises(ValueError): export_statement(out,q,d,matches,draft=False,acknowledged=True)
        edit_item(d,0,{'quantity':'2'});d=commit_review(d)
        self.assertFalse(d.forced)
        self.assertEqual(export_statement(out,q,d,reconcile(q.items,d.items),draft=False,acknowledged=True),D(12))
        wb=openpyxl.load_workbook(out,data_only=True)
        self.assertEqual(json.loads(wb['来源明细']['W2'].value)['quantity'],'两箱')
        self.assertEqual(json.loads(wb['来源明细']['X2'].value)[0]['after'],'2');wb.close()

    def test_structural_failures_require_force_and_cannot_be_cleared_by_row_edit(self):
        p=self.book('d.xlsx',['品名','数量'],[['螺丝',2]])
        wb=openpyxl.load_workbook(p);wb.create_sheet('未知单据')['A1']='不能识别的表';wb.save(p);wb.close()
        draft=import_excel(p,'delivery',preview=True)
        self.assertEqual(len(draft.blocking_issues),1)
        edit_item(draft,0,{'quantity':'3'})
        with self.assertRaises(ValueError):commit_review(draft)
        self.assertTrue(commit_review(draft,force=True).forced)
        empty=self.book('empty.xlsx',['别的列'],[['说明']])
        with self.assertRaises(ImportFailure):import_excel(empty,'delivery',preview=True)

    def test_dates_and_formula_without_cache_remain_editable(self):
        p=self.book('d.xlsx',['品名','数量','单价'],[['螺丝',2,'=2+3']],meta='日期：2026-02-30 送货单号：A1')
        d=import_excel(p,'delivery',preview=True)
        self.assertIn('day',d.items[0].validation_errors)
        self.assertEqual(d.items[0].order,'A1')
        self.assertIn('price',d.items[0].validation_errors)
        edit_item(d,0,{'day':'2026-02-28','price':'5'})
        self.assertEqual(commit_review(d).items[0].day,date(2026,2,28))
        p=self.book('q.xlsx',['品名','单价'],[['螺丝','oops'],['螺母',-1]],meta='报价日期：2026-02-30')
        q=import_excel(p,'quote',preview=True)
        self.assertEqual(len(q.items),2)
        for i in q.items:
            self.assertIn('price',i.validation_errors);self.assertIn('effective',i.validation_errors)

    def test_original_amount_must_be_corrected_and_returns_allowed(self):
        p=self.book('d.xlsx',['品名','数量','单价','金额'],[['螺丝',2,3,7]])
        d=import_excel(p,'delivery',preview=True)
        self.assertIn('original_amount',d.items[0].validation_errors)
        edit_item(d,0,{'quantity':'-2','original_amount':'-6'})
        self.assertEqual(commit_review(d).items[0].quantity,D(-2))

    def test_progress_restores_edits_invalid_raw_and_manual_decision(self):
        qp=self.book('q.xlsx',['品名','单价'],[['螺丝',3]])
        dp=self.book('d.xlsx',['品名','数量'],[['螺丝','2箱'],['螺丝','错']])
        q=commit_review(import_excel(qp,'quote',preview=True))
        d=import_excel(dp,'delivery',preview=True);edit_item(d,0,{'quantity':'2'})
        d=commit_review(d,force=True);matches=reconcile(q.items,d.items)
        confirm(matches[0],q.items[0],D(1),D(0),'赠送确认')
        path=self.root/'progress.json';save_session(path,q,d,matches)
        q2,d2,m2=load_session(path)
        self.assertEqual(d2.items[0].quantity,D(2))
        self.assertEqual(d2.items[0].original_values['quantity'],'2箱')
        self.assertEqual(d2.items[1].edit_values['quantity'],'错')
        self.assertEqual(d2.unresolved,1);self.assertTrue(d2.forced)
        self.assertEqual(m2[0].amount,D(0));self.assertFalse(m2[1].ready)


if __name__=='__main__': unittest.main()
