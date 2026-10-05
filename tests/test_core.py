"""合成数据的业务回归测试，无需真实业务文件或 Office。"""
import hashlib
from datetime import date
from decimal import Decimal as D
from pathlib import Path
import tempfile
import unittest
import openpyxl
from app.models import Item, ImportResult, ImportFailure
from app.parser import import_excel
from app.matching import reconcile, confirm
from app.exporter import export_statement, chinese_money
from app.session import save_session, load_session


class CoreTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory(); self.root=Path(self.tmp.name)

    def tearDown(self):self.tmp.cleanup()

    def workbook(self,name,headers,rows,start=4,col=3,meta=None):
        wb=openpyxl.Workbook();ws=wb.active;ws.title='明细'
        if meta:ws.cell(1,1,meta)
        for c,h in enumerate(headers,col):ws.cell(start,c,h)
        for r,values in enumerate(rows,start+1):
            for c,v in enumerate(values,col):ws.cell(r,c,v)
        # 验证巨大空白使用区域不会形成巨大矩阵。
        ws.cell(1000,16000).number_format='0.00'
        path=self.root/name;wb.save(path);return path

    def item(self,name='螺丝4*20',unit='斤',price='5.5',qty=None,serial='1',code=''):
        return Item('source.xlsx','明细',7,serial,name,unit=unit,price=D(price) if price is not None else None,quantity=D(qty) if qty else None,code=code)

    def test_offset_priority_empty_rows_and_missing_prices(self):
        p=self.workbook('q.xlsx',['序号','商品名称','单位','参考报价','圆养报价'],[['Q1','螺丝','斤',99,5.5],[None]*5,['Q2','待价','盒',8,None]])
        result=import_excel(p,'quote')
        self.assertEqual(len(result.items),2);self.assertEqual(result.items[0].price,D('5.5'))
        self.assertIsNone(result.items[1].price);self.assertIn('缺少有效报价',result.diagnostic())

    def test_name_and_spec_separate_matches_combined(self):
        q=self.item(name='螺丝');q.spec='4×20'
        d=self.item(name='螺丝4*20',qty='3')
        self.assertTrue(reconcile([q],[d])[0].ready)

    def test_normalization_duplicate_and_units(self):
        q=self.item(name='螺丝４ × ２０.０')
        d=self.item(name=' 螺丝4*20 ',qty='2')
        self.assertTrue(reconcile([q],[d])[0].ready)
        self.assertFalse(reconcile([q,q],[d])[0].ready)
        d.unit='箱';self.assertEqual(reconcile([q],[d])[0].status,'单位不一致')

    def test_code_conflict_and_serial_not_identity(self):
        q=self.item(code='SKU1');d=self.item(name='螺丝4*30',code='SKU1',qty='2')
        self.assertFalse(reconcile([q],[d])[0].ready)
        q.code=d.code='';self.assertFalse(reconcile([q],[d])[0].ready)

    def test_ambiguous_and_missing_headers(self):
        p=self.workbook('bad.xlsx',['品名','备注'],[['螺丝','abc']])
        with self.assertRaisesRegex(ImportFailure,'找不到单价列'):import_excel(p,'quote')
        p=self.workbook('amb.xlsx',['品名','单价','单价'],[['螺丝',1,2]])
        with self.assertRaisesRegex(ImportFailure,'多个列'):import_excel(p,'quote')
        p=self.workbook('layouts.xlsx',['品名','单价'],[['螺丝',1],['单价','品名'],[2,'螺母']])
        with self.assertRaisesRegex(ImportFailure,'多个不同布局'):import_excel(p,'quote')

    def test_empty_and_partial_rows(self):
        p=self.workbook('empty.xlsx',['序号','品名','单位','数量'],[[1,None,'条',None]])
        with self.assertRaisesRegex(ImportFailure,'数据区域为空'):import_excel(p,'delivery')
        p=self.workbook('partial.xlsx',['品名','数量'],[['螺丝',None]])
        with self.assertRaisesRegex(ImportFailure,'数量为空'):import_excel(p,'delivery')

    def test_unrecognized_sheet_never_silently_dropped(self):
        p=self.workbook('sheet.xlsx',['品名','数量'],[['螺丝',2]])
        wb=openpyxl.load_workbook(p);s=wb.create_sheet('另一个订单');s['A1']='商品资料';wb.save(p)
        with self.assertRaisesRegex(ImportFailure,'另一个订单'):import_excel(p,'delivery')

    def test_formula_without_cache_and_bad_number(self):
        p=self.workbook('formula.xlsx',['品名','单位','单价'],[['螺丝','斤','=2+3']])
        with self.assertRaisesRegex(ImportFailure,'公式没有已保存'):import_excel(p,'quote')
        p=self.workbook('number.xlsx',['品名','数量'],[['螺丝','2箱']])
        with self.assertRaisesRegex(ImportFailure,'不是有效数字'):import_excel(p,'delivery')

    def test_dates_meta_merged_and_formula_amount(self):
        p=self.workbook('delivery.xlsx',['序号','品名','数量','日期','送货单号','单价','金额'],[[1,'螺丝',2,date(2026,8,1),'S1',D('1.2'),5]])
        result=import_excel(p,'delivery');self.assertEqual(result.items[0].day,date(2026,8,1));self.assertIn('不一致',result.diagnostic())
        p=self.workbook('meta.xlsx',['品名','数量'],[['螺丝',2],[None,3]],meta='送货单号：S-20260801 日期：2026-08-01')
        wb=openpyxl.load_workbook(p);wb.active.merge_cells('C5:C6');wb.save(p)
        result=import_excel(p,'delivery');self.assertEqual(len(result.items),2);self.assertEqual(result.items[1].name,'螺丝');self.assertEqual(result.items[0].order,'S-20260801')

    def test_manual_conversion_gift_and_validation(self):
        q=self.item(unit='盒',price='8.2');d=self.item(unit='箱',qty='3')
        m=reconcile([q],[d])[0]
        with self.assertRaisesRegex(ValueError,'确认原因'):confirm(m,q,D(20),None,'')
        confirm(m,q,D(20),None,'一箱20盒，包装确认');self.assertEqual(m.amount,D(492))
        confirm(m,q,D(20),D(0),'赠送');self.assertEqual(m.amount,D(0))
        d.note='赠送';self.assertFalse(reconcile([q],[d])[0].ready)

    def test_export_formulas_provenance_tax_and_input_unchanged(self):
        qp=self.workbook('q.xlsx',['序号','品名','单位','单价'],[['Q1','零件','个',.283]])
        dp=self.workbook('d.xlsx',['序号','品名','单位','数量'],[['D9','零件','个',3728]])
        before=[hashlib.sha256(p.read_bytes()).hexdigest() for p in [qp,dp]]
        q=import_excel(qp,'quote');d=import_excel(dp,'delivery');m=reconcile(q.items,d.items)
        out=self.root/'out.xlsx'
        with self.assertRaisesRegex(ValueError,'正式对账单要求'):export_statement(out,q,d,m,draft=False)
        total=export_statement(out,q,d,m,draft=False,acknowledged=True,tax_rate=D('.06'))
        self.assertEqual(total,D('1055.024'))
        wb=openpyxl.load_workbook(out);ws=wb.active
        self.assertEqual(ws['H7'].value,'=F7*G7');self.assertEqual(ws['J7'].value,'Q1');self.assertEqual(ws['K7'].value,'D9')
        cached=openpyxl.load_workbook(out,data_only=True).active
        self.assertAlmostEqual(cached['H7'].value,1055.024);self.assertEqual(cached['H10'].value,63.30)
        self.assertIn('A1:K1',str(ws.merged_cells));self.assertEqual(len(wb.sheetnames),2)
        self.assertEqual(before,[hashlib.sha256(p.read_bytes()).hexdigest() for p in [qp,dp]])
        with self.assertRaisesRegex(ValueError,'不能覆盖'):export_statement(qp,q,d,m)

    def test_draft_excludes_unresolved_and_literal_names(self):
        q=ImportResult(str(self.root/'q.xlsx'),[self.item(name='=HYPERLINK("x")')])
        d=ImportResult(str(self.root/'d.xlsx'),[self.item(name='=HYPERLINK("x")',qty='2'),self.item(name='未知商品',qty='100')])
        m=reconcile(q.items,d.items);out=self.root/'draft.xlsx'
        export_statement(out,q,d,m)
        ws=openpyxl.load_workbook(out).active
        self.assertEqual(ws['D7'].data_type,'s');self.assertIsNone(ws['H8'].value)
        self.assertIn('待核对稿',ws['A2'].value);self.assertIn('已确认部分小计',ws['A10'].value)
        with self.assertRaises(ValueError):export_statement(out,q,d,m,draft=False,acknowledged=True)

    def test_tax_included_no_double_charge_and_negative_quantity(self):
        qi=self.item();qi.tax_included=True
        q=ImportResult(str(self.root/'q.xlsx'),[qi]);d=ImportResult(str(self.root/'d.xlsx'),[self.item(qty='-2')])
        m=reconcile(q.items,d.items);out=self.root/'return.xlsx'
        self.assertEqual(export_statement(out,q,d,m,tax_rate=D('.06')),D('-11'))
        self.assertEqual(openpyxl.load_workbook(out,data_only=True).active['H10'].value,0)

    def test_progress_rejects_changed_sources(self):
        qp=self.workbook('q.xlsx',['品名','单价'],[['螺丝',5]])
        dp=self.workbook('d.xlsx',['品名','数量'],[['螺丝',2]])
        q=import_excel(qp,'quote');d=import_excel(dp,'delivery');m=reconcile(q.items,d.items)
        confirm(m[0],q.items[0],D(1),D(0),'赠送确认')
        progress=self.root/'session.json';save_session(progress,q,d,m)
        q2,d2,m2=load_session(progress);self.assertEqual(m2[0].amount,D(0));self.assertEqual(m2[0].reason,'赠送确认')
        with qp.open('ab') as f:f.write(b'changed')
        with self.assertRaisesRegex(ValueError,'来源文件已变化'):load_session(progress)

    def test_money_uppercase(self):
        self.assertEqual(chinese_money(D('1044.9')),'壹仟零肆拾肆元玖角')
        self.assertEqual(chinese_money(D('10001.01')),'壹万零壹元零壹分')
        self.assertEqual(chinese_money(D('0')),'零元整')

    def test_multiple_customers_rejected(self):
        p=self.workbook('customers.xlsx',['品名','数量'],[['螺丝',2]],meta='TO：客户甲')
        wb=openpyxl.load_workbook(p);ws=wb.copy_worksheet(wb.active);ws.title='第二单';ws['A1']='TO：客户乙';wb.save(p)
        with self.assertRaisesRegex(ImportFailure,'不同客户'):import_excel(p,'delivery')

    def test_repeated_delivery_blocks(self):
        p=self.workbook('blocks.xlsx',['品名','数量'],[['螺丝',2],['合计',2],['送货单号：B2 日期：2026-10-02',None],['品名','数量'],['螺母',3]],meta='送货单号：B1 日期：2026-10-01')
        result=import_excel(p,'delivery')
        self.assertEqual([i.order for i in result.items],['B1','B2'])
        self.assertEqual([i.day for i in result.items],[date(2026,10,1),date(2026,10,2)])

if __name__=='__main__':unittest.main()
