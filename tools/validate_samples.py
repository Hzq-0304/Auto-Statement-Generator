"""独立核验导出记录及样例差异，详细报告仅保存在被 Git 忽略的 analysis。"""
from pathlib import Path
from decimal import Decimal
from collections import Counter
import json
import sys
import hashlib
import openpyxl
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from app.parser import import_excel
from app.matching import reconcile
from app.exporter import export_statement

root=Path(__file__).resolve().parents[1]
qp=next(root.glob('*报价*.xlsx'));dp=next(root.glob('送货*.xlsx'));sp=next(root.glob('*对账单.xlsx'))
before={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in [qp,dp,sp]}
q=import_excel(qp,'quote');d=import_excel(dp,'delivery');matches=reconcile(q.items,d.items)
output=root/'outputs/sample/自动对账_待核对稿.xlsx'
total=export_statement(output,q,d,matches)
wb=openpyxl.load_workbook(output,data_only=True);ws=wb.active;audit=wb['来源明细']
for index,m in enumerate(matches,7):
    item=m.delivery
    assert ws.cell(index,11).value==item.serial
    assert ws.cell(index,3).value==item.order
    assert ws.cell(index,6).value==float(item.quantity)
    assert audit.cell(index-5,13).value==item.sheet
    assert audit.cell(index-5,14).value==item.row
    if m.ready:
        assert abs(Decimal(str(ws.cell(index,8).value))-m.amount)<Decimal('.00000001')
        assert ws.cell(index,10).value==m.quote.serial
    else:assert ws.cell(index,8).value is None
reference=openpyxl.load_workbook(sp,data_only=True)['202608']
historical=[r for r in reference.iter_rows(min_row=7,max_row=108) if isinstance(r[0].value,(int,float))]
assert len(historical)==102
assert [ws.cell(6,c).value for c in range(1,10)]==[reference.cell(6,c).value for c in range(1,10)]
historical_pairs=Counter((str(r[3].value).strip(),str(r[5].value)) for r in historical)
delivery_pairs=Counter((i.name.strip(),str(i.quantity)) for i in d.items)
report={'quotation_items':len(q.items),'missing_prices':sum(i.price is None for i in q.items),
        'delivery_items':len(d.items),'historical_statement_items':len(historical),
        'matched':sum(m.ready for m in matches),'pending':sum(not m.ready for m in matches),
        'matched_subtotal':str(total),'statuses':dict(Counter(m.status for m in matches)),
        'historical_total':str(reference['H110'].value),'historical_adjusted':str(reference['I110'].value),
        'common_exact_name_quantity_count':sum((historical_pairs & delivery_pairs).values()),
        'source_sha256':before,'all_99_output_rows_and_sources_checked':True,'original_A_to_I_headers_preserved':True}
assert before=={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in [qp,dp,sp]}
(root/'analysis/validation-report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in report.items() if k!='source_sha256'},ensure_ascii=False,indent=2))
